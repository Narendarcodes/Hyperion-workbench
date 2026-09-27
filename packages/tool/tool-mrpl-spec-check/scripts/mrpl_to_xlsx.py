#!/usr/bin/env python3
"""Deterministic MRPL MG 91 compliance workbook builder.

Reads a mrpl_spec_check result object (stdin JSON) plus approval metadata and
writes S-001_MG91_Compliance.xlsx. No LLM, no network. The workbook dataset
comes ONLY from the checker output; this script adds layout, the demo-input
banner, and provenance. Requires openpyxl (pip install openpyxl once).

Usage:
    py mrpl_to_xlsx.py --out S-001_MG91_Compliance.xlsx < payload.json
    payload = {"check_result": {...}, "approved_by": "...",
               "approval_note": "...", "approval_time": "..."}
Prints {"out_path": ..., "overall": ..., "rows": N} on stdout. Exit 2 on error.
"""

import datetime
import json
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    print(json.dumps({"code": "missing_dependency",
                      "errors": ["openpyxl is not installed (py -m pip install openpyxl)"]}))
    sys.exit(2)

BANNER = "Sample %s - demo input, not a refinery certificate."
RESULT_HEADERS = ["Sample ID", "Parameter", "Specification", "Lab Value",
                  "Unit", "Margin", "Verdict", "Method", "Citation"]
FAIL_FILL = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
PASS_FILL = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
NA_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
BOLD = Font(bold=True)


def fail(code, errors):
    print(json.dumps({"code": code, "errors": errors}))
    sys.exit(2)


def main(argv):
    out_path = None
    idx = 0
    while idx < len(argv):
        if argv[idx] == "--out":
            idx += 1
            out_path = argv[idx] if idx < len(argv) else None
        idx += 1
    if not out_path:
        fail("bad_request", ["--out PATH is required"])
    try:
        payload = json.load(sys.stdin)
    except (json.JSONDecodeError, UnicodeDecodeError) as exc:
        fail("bad_request", ["payload unreadable: %s" % exc])
    result = payload.get("check_result", {})
    rows = result.get("rows", [])
    if not isinstance(rows, list) or not rows:
        fail("bad_request", ["payload check_result.rows must be a non-empty list"])
    approver = str(payload.get("approved_by", "") or "").strip()
    if not approver:
        fail("approval_required", ["approved_by is required"])
    sample_id = str(result.get("sample_id", "S-001"))
    approval_time = str(payload.get("approval_time", "") or "") \
        or datetime.datetime.now().isoformat(timespec="seconds")

    book = Workbook()
    sheet = book.active
    sheet.title = "Results"
    sheet.merge_cells("A1:I1")
    sheet["A1"] = BANNER % sample_id
    sheet["A1"].font = BOLD
    for col, header in enumerate(RESULT_HEADERS, start=1):
        cell = sheet.cell(row=2, column=col, value=header)
        cell.font = BOLD
    for offset, item in enumerate(rows):
        line = 3 + offset
        values = [sample_id, item.get("parameter"), item.get("specification"),
                  item.get("lab_value"), item.get("unit"), item.get("margin"),
                  item.get("verdict"), item.get("method"), item.get("citation")]
        for col, value in enumerate(values, start=1):
            sheet.cell(row=line, column=col, value=value)
        verdict = item.get("verdict")
        fill = FAIL_FILL if verdict == "FAIL" else PASS_FILL if verdict == "PASS" else NA_FILL
        for col in range(1, 10):
            sheet.cell(row=line, column=col).fill = fill
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = "A2:I%d" % (2 + len(rows))
    for col, width in zip("ABCDEFGHI", (12, 24, 18, 12, 10, 10, 10, 44, 14)):
        sheet.column_dimensions[col].width = width

    prov = book.create_sheet("Provenance")
    prov["A1"] = BANNER % sample_id
    prov["A1"].font = BOLD
    provenance = result.get("provenance", {})
    summary = result.get("summary", {})
    entries = [
        ("Sample ID", sample_id),
        ("Spec ID", result.get("spec_id", "DOC-010")),
        ("Spec title", result.get("spec_title", "")),
        ("Season", result.get("season", "")),
        ("Overall", summary.get("overall", "")),
        ("Pass / Fail / NA", "%s / %s / %s" % (summary.get("pass"), summary.get("fail"), summary.get("na"))),
        ("Source", provenance.get("source", "")),
        ("Source hash", provenance.get("source_hash", "")),
        ("Tool version", provenance.get("tool_version", "")),
        ("Generated at", datetime.datetime.now().isoformat(timespec="seconds")),
        ("Approval status", "APPROVED"),
        ("Approver", approver),
        ("Approval time", approval_time),
        ("Approval note", str(payload.get("approval_note", "") or "")),
    ]
    prov.cell(row=2, column=1, value="Field").font = BOLD
    prov.cell(row=2, column=2, value="Value").font = BOLD
    for offset, (field, value) in enumerate(entries):
        prov.cell(row=3 + offset, column=1, value=field)
        prov.cell(row=3 + offset, column=2, value=value)
    prov.column_dimensions["A"].width = 18
    prov.column_dimensions["B"].width = 90
    for row in prov.iter_rows(min_row=1, max_row=2 + len(entries), max_col=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    try:
        book.save(out_path)
    except OSError as exc:
        fail("write_failed", ["cannot write %s: %s" % (out_path, exc)])
    print(json.dumps({"out_path": out_path,
                      "overall": summary.get("overall", ""),
                      "rows": len(rows)}))


if __name__ == "__main__":
    main(sys.argv[1:])
